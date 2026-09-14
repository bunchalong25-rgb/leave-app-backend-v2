import os
import json
import uuid
import io
import urllib.parse
from datetime import datetime
from typing import Optional, List, Dict
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Response, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import RedirectResponse
from pydantic import BaseModel
import openpyxl
import httpx
from excel_processor import create_sample_template, process_excel_template

app = FastAPI(title="Staff Leave Management API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

DB_FILE = os.path.join(os.path.dirname(__file__), "db.json")
UPLOADS_DIR = os.path.join(os.path.dirname(__file__), "uploads")
os.makedirs(UPLOADS_DIR, exist_ok=True)
SAMPLE_TEMPLATE_PATH = os.path.join(UPLOADS_DIR, "sample_template.xlsx")

if not os.path.exists(SAMPLE_TEMPLATE_PATH):
    create_sample_template(SAMPLE_TEMPLATE_PATH)

# Clean initial database seed
INITIAL_DB = {
    "brands": [],
    "employees": [],
    "users": [
        {
            "id": "usr_admin",
            "lineUserId": "U_ADMIN_DEMO",
            "displayName": "ดารินทร์ สุขใจ (Admin)",
            "fullName": "ดารินทร์ สุขใจ",
            "brand": "สำนักงานใหญ่",
            "role": "admin",
            "hasCompletedOnboarding": True,
            "createdAt": datetime.utcnow().isoformat()
        }
    ],
    "cutoffs": [
        {
            "id": "cut_2026_09_1",
            "yearMonth": "2026-09",
            "period": "1-15",
            "cutoffDatetime": "2026-09-25T23:59:59.000Z",
            "updatedAt": datetime.utcnow().isoformat()
        },
        {
            "id": "cut_2026_09_2",
            "yearMonth": "2026-09",
            "period": "16-end",
            "cutoffDatetime": "2026-09-30T23:59:59.000Z",
            "updatedAt": datetime.utcnow().isoformat()
        }
    ],
    "leaveRecords": []
}

def read_db():
    if not os.path.exists(DB_FILE):
        write_db(INITIAL_DB)
        return INITIAL_DB
    try:
        with open(DB_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        write_db(INITIAL_DB)
        return INITIAL_DB

def write_db(data):
    with open(DB_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

# Request Models
class LineLoginRequest(BaseModel):
    lineUserId: str
    displayName: Optional[str] = None

class OnboardingRequest(BaseModel):
    lineUserId: str
    displayName: Optional[str] = None
    fullName: str
    brand: str
    role: str # 'staff' or 'admin'
    supervisorId: Optional[str] = None
    supervisorName: Optional[str] = None
    inviteCode: Optional[str] = None

class CutoffRequest(BaseModel):
    yearMonth: str
    period: str
    cutoffDatetime: str

class LeaveSubmitRequest(BaseModel):
    userId: str
    userName: str
    brand: str
    yearMonth: str
    period: str
    dateCodeMap: Dict[str, Optional[str]]

class ExcelProcessRequest(BaseModel):
    yearMonth: str
    period: Optional[str] = "1-15"

class BrandRequest(BaseModel):
    name: str

class EmployeeRequest(BaseModel):
    name: str
    brand: str

# Base Configs
FRONTEND_URL = os.getenv("FRONTEND_URL", "https://creative-macaron-98004d.netlify.app")
LINE_CHANNEL_ID = os.getenv("LINE_CHANNEL_ID", "")
LINE_CHANNEL_SECRET = os.getenv("LINE_CHANNEL_SECRET", "")
LINE_CALLBACK_URL = os.getenv("LINE_CALLBACK_URL", "https://leave-app-backend-v2.onrender.com/api/auth/line/callback")

# Routes
@app.get("/")
def read_root():
    return {"message": "Staff Leave Management API is online", "status": "ok"}

# ---------------------------------------------------------
# REAL LINE LOGIN OAUTH 2.0 FLOW
# ---------------------------------------------------------

@app.get("/api/auth/line/login")
def line_oauth_login():
    channel_id = LINE_CHANNEL_ID
    if not channel_id:
        # Fallback if channel ID env variable not yet configured
        return RedirectResponse(url=f"{FRONTEND_URL}/?error=LINE_CHANNEL_ID_NOT_CONFIGURED")

    state = uuid.uuid4().hex[:12]
    params = {
        "response_type": "code",
        "client_id": channel_id,
        "redirect_uri": LINE_CALLBACK_URL,
        "state": state,
        "scope": "profile openid"
    }

    line_authorize_url = f"https://access.line.me/oauth2/v2.1/authorize?{urllib.parse.urlencode(params)}"
    return RedirectResponse(url=line_authorize_url)


@app.get("/api/auth/line/callback")
async def line_oauth_callback(code: Optional[str] = None, state: Optional[str] = None, error: Optional[str] = None):
    if error or not code:
        return RedirectResponse(url=f"{FRONTEND_URL}/?error={error or 'NO_CODE'}")

    try:
        # 1. Exchange Auth Code for Access Token
        token_url = "https://api.line.me/oauth2/v2.1/token"
        token_data = {
            "grant_type": "authorization_code",
            "code": code,
            "redirect_uri": LINE_CALLBACK_URL,
            "client_id": LINE_CHANNEL_ID,
            "client_secret": LINE_CHANNEL_SECRET,
        }

        async with httpx.AsyncClient() as client:
            token_res = await client.post(
                token_url,
                data=token_data,
                headers={"Content-Type": "application/x-www-form-urlencoded"}
            )

            if token_res.status_code != 200:
                print("LINE Token exchange failed:", token_res.text)
                return RedirectResponse(url=f"{FRONTEND_URL}/?error=TOKEN_EXCHANGE_FAILED")

            tokens = token_res.json()
            access_token = tokens.get("access_token")

            # 2. Get User Profile from LINE API
            profile_res = await client.get(
                "https://api.line.me/v2/profile",
                headers={"Authorization": f"Bearer {access_token}"}
            )

            if profile_res.status_code != 200:
                print("LINE Profile fetch failed:", profile_res.text)
                return RedirectResponse(url=f"{FRONTEND_URL}/?error=PROFILE_FETCH_FAILED")

            profile = profile_res.json()
            line_user_id = profile.get("userId")
            display_name = profile.get("displayName", "")
            picture_url = profile.get("pictureUrl", "")

            # 3. Check DB if User exists
            db = read_db()
            users = db.get("users", [])
            existing_user = next((u for u in users if u.get("lineUserId") == line_user_id), None)

            # 4. Redirect back to Frontend Netlify with User info Query Params
            if existing_user:
                user_json = urllib.parse.quote(json.dumps(existing_user))
                return RedirectResponse(url=f"{FRONTEND_URL}/?exists=true&user={user_json}")
            else:
                encoded_name = urllib.parse.quote(display_name)
                encoded_pic = urllib.parse.quote(picture_url)
                return RedirectResponse(
                    url=f"{FRONTEND_URL}/?exists=false&lineUserId={line_user_id}&displayName={encoded_name}&pictureUrl={encoded_pic}"
                )

    except Exception as e:
        print("Error in LINE OAuth Callback:", str(e))
        return RedirectResponse(url=f"{FRONTEND_URL}/?error=INTERNAL_CALLBACK_ERROR")


# ---------------------------------------------------------
# AUTHENTICATION & ONBOARDING API
# ---------------------------------------------------------

@app.post("/api/auth/login-line")
def login_line(req: LineLoginRequest):
    db = read_db()
    users = db.get("users", [])
    matched = next((u for u in users if u.get("lineUserId") == req.lineUserId), None)
    if matched:
        return {"exists": True, "user": matched}
    else:
        return {"exists": False, "lineUserId": req.lineUserId, "displayName": req.displayName}

@app.get("/api/onboarding/options")
def get_onboarding_options():
    db = read_db()
    users = db.get("users", [])
    supervisors = [{"id": u["id"], "fullName": u["fullName"]} for u in users if u.get("role") == "admin"]
    brands = db.get("brands", [])
    employees = db.get("employees", [])
    
    claimed_names = set(u.get("fullName") for u in users if u.get("role") == "staff")
    available_employees = [e for e in employees if e.get("name") not in claimed_names]

    return {
        "supervisors": supervisors,
        "brands": brands,
        "masterEmployees": available_employees,
        "hasMasterData": len(employees) > 0 and len(supervisors) > 0
    }

@app.post("/api/auth/register-onboarding")
def register_onboarding(req: OnboardingRequest):
    if req.role == "admin":
        invite_code_env = os.getenv("ADMIN_INVITE_CODE", "ADMIN2026")
        if not req.inviteCode or req.inviteCode.strip().upper() != invite_code_env:
            raise HTTPException(status_code=403, detail="รหัสผ่านพิเศษ (Invite Code) สำหรับหัวหน้าแผนกไม่ถูกต้อง")

    db = read_db()
    user = {
        "id": f"usr_{int(datetime.utcnow().timestamp())}",
        "lineUserId": req.lineUserId,
        "displayName": req.displayName or req.fullName,
        "fullName": req.fullName,
        "brand": req.brand,
        "role": req.role,
        "supervisorId": req.supervisorId,
        "supervisorName": req.supervisorName,
        "hasCompletedOnboarding": True,
        "createdAt": datetime.utcnow().isoformat()
    }
    db["users"].append(user)
    write_db(db)
    return {"success": True, "user": user}

# ---------------------------------------------------------
# ADMIN MASTER DATA UPLOAD & STATUS DASHBOARD
# ---------------------------------------------------------

@app.post("/api/admin/upload-master-data")
async def upload_master_data(file: UploadFile = File(...)):
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="กรุณาอัปโหลดไฟล์ Excel (.xlsx)")

    content = await file.read()
    wb = openpyxl.load_workbook(filename=io.BytesIO(content))
    ws = wb.active

    db = read_db()
    existing_employees = db.get("employees", [])
    existing_brands = db.get("brands", [])

    added_emp_count = 0
    added_brand_count = 0

    name_col = 1
    brand_col = 2

    for c in range(1, ws.max_column + 1):
        val = str(ws.cell(row=1, column=c).value or "").strip().lower()
        if "ชื่อ" in val or "name" in val:
            name_col = c
        elif "แบรนด์" in val or "brand" in val:
            brand_col = c

    for r in range(2, ws.max_row + 1):
        emp_name = str(ws.cell(row=r, column=name_col).value or "").strip()
        brand_name = str(ws.cell(row=r, column=brand_col).value or "").strip()

        if emp_name:
            if not any(e["name"] == emp_name for e in existing_employees):
                existing_employees.append({
                    "id": f"emp_{int(datetime.utcnow().timestamp())}_{r}",
                    "name": emp_name,
                    "brand": brand_name or "General"
                })
                added_emp_count += 1

        if brand_name:
            if not any(b["name"] == brand_name for b in existing_brands):
                existing_brands.append({
                    "id": f"b_{int(datetime.utcnow().timestamp())}_{r}",
                    "name": brand_name
                })
                added_brand_count += 1

    db["employees"] = existing_employees
    db["brands"] = existing_brands
    write_db(db)

    return {
        "success": True,
        "message": f"นำเข้าข้อมูลสำเร็จ: เพิ่มพนักงาน {added_emp_count} คน, เพิ่มแบรนด์ {added_brand_count} แบรนด์",
        "addedEmployees": added_emp_count,
        "addedBrands": added_brand_count
    }

@app.get("/api/admin/submission-status")
def get_submission_status(yearMonth: str = "2026-09", period: str = "1-15", supervisorId: Optional[str] = None):
    db = read_db()
    users = db.get("users", [])
    leave_records = db.get("leaveRecords", [])

    staff_users = [u for u in users if u.get("role") == "staff"]
    if supervisorId:
        staff_users = [u for u in staff_users if u.get("supervisorId") == supervisorId]

    submitted_user_ids = set(
        r.get("userId") for r in leave_records
        if r.get("yearMonth") == yearMonth and r.get("period") == period
    )

    submitted_list = []
    pending_list = []

    for u in staff_users:
        info = {
            "userId": u.get("id"),
            "fullName": u.get("fullName"),
            "brand": u.get("brand"),
            "supervisorName": u.get("supervisorName", "-")
        }
        if u.get("id") in submitted_user_ids:
            submitted_list.append(info)
        else:
            pending_list.append(info)

    return {
        "totalStaff": len(staff_users),
        "submittedCount": len(submitted_list),
        "pendingCount": len(pending_list),
        "submittedList": submitted_list,
        "pendingList": pending_list
    }

# ---------------------------------------------------------
# CUTOFFS & LEAVE RECORDS
# ---------------------------------------------------------

@app.get("/api/brands")
def get_brands():
    db = read_db()
    return {"brands": db.get("brands", [])}

@app.get("/api/employees")
def get_employees():
    db = read_db()
    return {"employees": db.get("employees", [])}

@app.get("/api/cutoffs")
def get_cutoffs():
    db = read_db()
    return {"cutoffs": db.get("cutoffs", [])}

@app.post("/api/cutoffs")
def set_cutoff(req: CutoffRequest):
    db = read_db()
    cutoffs = db.get("cutoffs", [])
    found = False
    for c in cutoffs:
        if c.get("yearMonth") == req.yearMonth and c.get("period") == req.period:
            c["cutoffDatetime"] = req.cutoffDatetime
            c["updatedAt"] = datetime.utcnow().isoformat()
            found = True
            break
    if not found:
        cutoffs.append({
            "id": f"cut_{int(datetime.utcnow().timestamp())}",
            "yearMonth": req.yearMonth,
            "period": req.period,
            "cutoffDatetime": req.cutoffDatetime,
            "updatedAt": datetime.utcnow().isoformat()
        })
    db["cutoffs"] = cutoffs
    write_db(db)
    return {"success": True, "cutoffs": cutoffs}

@app.get("/api/leaves")
def get_leaves(yearMonth: str = "2026-09", period: Optional[str] = None):
    db = read_db()
    records = db.get("leaveRecords", [])
    filtered = [r for r in records if r.get("yearMonth") == yearMonth and (not period or r.get("period") == period)]
    return {"records": filtered}

@app.post("/api/leaves")
def save_leaves(req: LeaveSubmitRequest):
    db = read_db()
    cutoffs = db.get("cutoffs", [])
    current_cutoff = next((c for c in cutoffs if c.get("yearMonth") == req.yearMonth and c.get("period") == req.period), None)

    if current_cutoff:
        cutoff_dt = datetime.fromisoformat(current_cutoff["cutoffDatetime"].replace("Z", "+00:00"))
        now_dt = datetime.utcnow()
        if now_dt > cutoff_dt.replace(tzinfo=None):
            raise HTTPException(status_code=403, detail="ระบบปิดรับการบันทึกข้อมูลสำหรับรอบนี้แล้ว ไม่สามารถแก้ไขได้")

    records = [
        r for r in db.get("leaveRecords", [])
        if not (r.get("userId") == req.userId and r.get("yearMonth") == req.yearMonth and r.get("period") == req.period)
    ]

    for date_str, code in req.dateCodeMap.items():
        if code:
            day_num = int(date_str.split("-")[2])
            records.append({
                "id": f"lr_{int(datetime.utcnow().timestamp())}_{uuid.uuid4().hex[:4]}",
                "userId": req.userId,
                "userName": req.userName,
                "brand": req.brand,
                "yearMonth": req.yearMonth,
                "period": req.period,
                "date": date_str,
                "dayNumber": day_num,
                "code": code,
                "updatedAt": datetime.utcnow().isoformat()
            })

    db["leaveRecords"] = records
    write_db(db)
    return {"success": True, "message": "บันทึกตารางวันหยุดสำเร็จ"}

active_template_file = SAMPLE_TEMPLATE_PATH

@app.post("/api/excel/upload-template")
async def upload_template(templateFile: UploadFile = File(...)):
    global active_template_file
    file_path = os.path.join(UPLOADS_DIR, f"custom_{templateFile.filename}")
    content = await templateFile.read()
    with open(file_path, "wb") as f:
        f.write(content)
    active_template_file = file_path
    return {"success": True, "message": "อัปโหลดไฟล์ Excel Template สำเร็จ", "filename": templateFile.filename}

@app.post("/api/excel/process")
def process_excel(req: ExcelProcessRequest):
    global active_template_file
    db = read_db()
    records = [
        r for r in db.get("leaveRecords", [])
        if r.get("yearMonth") == req.yearMonth and (not req.period or r.get("period") == req.period)
    ]

    try:
        excel_bytes, updated_count = process_excel_template(
            template_path=active_template_file,
            leave_records=records,
            period=req.period
        )

        filename = f"Leave_Schedule_{req.yearMonth}_{req.period or 'All'}.xlsx"
        headers = {
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-Updated-Cells-Count": str(updated_count),
            "Access-Control-Expose-Headers": "Content-Disposition, X-Updated-Cells-Count"
        }

        return Response(
            content=excel_bytes,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers=headers
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
